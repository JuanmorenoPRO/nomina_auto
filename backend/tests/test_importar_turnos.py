"""Importación de turnos desde la plantilla de Excel (formato tarjeta).

El test que más protege es el de ida y vuelta: se descarga la plantilla con los
turnos que hay en la base y se vuelve a subir tal cual; la quincena tiene que
quedar idéntica. Si un cambio rompe el formato en cualquiera de las dos
direcciones, este test lo ve.
"""

from __future__ import annotations

from datetime import datetime, time
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --- Armado de archivos de prueba (no usa la plantilla real: así los tests
# fijan el contrato del lector, no la implementación del generador) ---


def construir_libro(hojas: list[dict[str, Any]], con_instrucciones: bool = True) -> bytes:
    """`hojas`: [{"nombre", "documento", "filas": [(día, entra, sale, jorn)],
    "marcas": {...}, "sin_cc": bool}]"""
    libro = Workbook()
    primera = libro.active
    assert primera is not None
    if con_instrucciones:
        primera.title = "INSTRUCCIONES"
        primera["A1"] = "Cómo llenar esta plantilla"
        primera["A2"] = "Hay una hoja por empleado."
    else:
        libro.remove(primera)

    for hoja_datos in hojas:
        hoja = libro.create_sheet(hoja_datos["nombre"][:31])
        hoja["A1"] = "EDIFICIO:"
        hoja["B1"] = hoja_datos.get("unidad", "EDIFICIO DE PRUEBA")
        hoja["A2"] = "NOMBRE Y APELLIDO:"
        hoja["B2"] = hoja_datos["nombre"]
        if not hoja_datos.get("sin_cc"):
            hoja["A3"] = "C.C.:"
            hoja["B3"] = hoja_datos["documento"]
        hoja["A4"] = "MES:"
        hoja["B4"] = "01/09/2026 al 15/09/2026"

        fila = 6
        etiquetas = {
            "quincena_incompleta": "QUINCENA INCOMPLETA:",
            "sin_extras": "NO CALCULAR EXTRAS:",
            "auxilio_por_dias_laborados": "AUXILIO PRORRATEADO:",
            "pagar_dia_31": "PAGAR DÍA 31:",
        }
        for clave, valor in (hoja_datos.get("marcas") or {}).items():
            hoja.cell(row=fila, column=1, value=etiquetas[clave])
            hoja.cell(row=fila, column=2, value="SÍ" if valor else "NO")
            fila += 1

        fila += 1
        for columna, texto in enumerate(
            ["DÍA", "N.º", "ENTRA", "SALE", "JORN. ORD. (h)", "TOTAL H"], start=1
        ):
            hoja.cell(row=fila, column=columna, value=texto)
        for dia, entra, sale, jornada in hoja_datos.get("filas", []):
            fila += 1
            hoja.cell(row=fila, column=1, value="lun")
            hoja.cell(row=fila, column=2, value=dia)
            hoja.cell(row=fila, column=3, value=entra)
            hoja.cell(row=fila, column=4, value=sale)
            hoja.cell(row=fila, column=5, value=jornada)

    contenido = BytesIO()
    libro.save(contenido)
    return contenido.getvalue()


def subir(client, periodo_id, unidad_id, contenido: bytes, validar_solo: bool = False):
    return client.post(
        f"/periodos/{periodo_id}/turnos/importar"
        f"?unidad_id={unidad_id}&validar_solo={'true' if validar_solo else 'false'}",
        files={"archivo": ("turnos.xlsx", contenido, XLSX)},
    )


def turnos_de(client, periodo_id, unidad_id, empleado_id=None):
    turnos = client.get(f"/periodos/{periodo_id}/turnos?unidad_id={unidad_id}").json()
    if empleado_id is not None:
        turnos = [t for t in turnos if t["empleado_id"] == empleado_id]
    return sorted(
        (
            (t["fecha"], t["hora_inicio"], t["hora_fin"], t["minutos_jornada_ordinaria"])
            for t in turnos
        )
    )


# --- Escenario base: una unidad con dos empleados y la quincena 1–15 sep 2026 ---


@pytest.fixture
def escenario(client):
    unidad = client.post("/unidades", json={"nombre": "EDIFICIO DE PRUEBA"}).json()
    empleados = []
    for nombre, documento in (("ANA GOMEZ", "1032456"), ("LUIS RIOS", "9988771")):
        r = client.post(
            "/empleados",
            json={
                "unidad_id": unidad["id"],
                "nombre": nombre,
                "documento": documento,
                "cargo": "vigilante",
                "salario_base": 1_750_905,
            },
        )
        assert r.status_code == 201, r.text
        empleados.append(r.json())
    periodo = client.post(
        "/periodos", json={"fecha_inicio": "2026-09-01", "fecha_fin": "2026-09-15"}
    ).json()
    return {"unidad": unidad, "empleados": empleados, "periodo": periodo}


def test_plantilla_y_reimportacion_dejan_los_mismos_turnos(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    # Turnos con todos los casos difíciles: cruce de medianoche, turno partido,
    # media hora y un turno de relleno (jornada ordinaria).
    for datos in (
        {"fecha": "2026-09-01", "hora_inicio": "06:00", "hora_fin": "18:00"},
        {"fecha": "2026-09-02", "hora_inicio": "18:00", "hora_fin": "06:00"},
        {"fecha": "2026-09-04", "hora_inicio": "06:00", "hora_fin": "10:00"},
        {"fecha": "2026-09-04", "hora_inicio": "14:30", "hora_fin": "18:00"},
        {
            "fecha": "2026-09-06",
            "hora_inicio": "06:00",
            "hora_fin": "13:00",
            "minutos_jornada_ordinaria": 420,
        },
    ):
        r = client.post("/turnos", json={"empleado_id": ana["id"], **datos})
        assert r.status_code == 201, r.text

    antes = turnos_de(client, periodo["id"], unidad["id"])
    assert len(antes) == 5

    r = client.get(f"/periodos/{periodo['id']}/turnos/plantilla?unidad_id={unidad['id']}")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    plantilla = r.content

    r = subir(client, periodo["id"], unidad["id"], plantilla)
    assert r.status_code == 200, r.text
    reporte = r.json()
    assert reporte["aplicado"] is True
    assert reporte["errores"] == []
    assert reporte["ignoradas"] == ["INSTRUCCIONES"]
    por_hoja = {h["empleado"]: h for h in reporte["hojas"]}
    assert por_hoja["ANA GOMEZ"]["turnos_creados"] == 5
    assert por_hoja["ANA GOMEZ"]["turnos_borrados"] == 5
    assert por_hoja["ANA GOMEZ"]["horas"] == "38.50"  # 12+12+4+3.5+7
    assert por_hoja["LUIS RIOS"]["turnos_creados"] == 0

    assert turnos_de(client, periodo["id"], unidad["id"]) == antes


def test_formatos_de_hora_que_devuelve_excel(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [
                    (1, time(6, 0), 18, None),  # hora de Excel + número entero
                    (2, "6,5", datetime(2026, 9, 2, 14, 0), None),  # coma decimal + datetime
                    (3, 18.5, "6:00", None),  # número con fracción + texto
                ],
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    assert r.json()["errores"] == []
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == [
        ("2026-09-01", "06:00:00", "18:00:00", None),
        ("2026-09-02", "06:30:00", "14:00:00", None),
        ("2026-09-03", "18:30:00", "06:00:00", None),
    ]


def test_turno_partido_y_jornada_ordinaria(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [
                    (4, "06:00", "10:00", None),
                    (4, "14:00", "18:00", None),  # misma fila de día: turno partido
                    (6, "06:00", "13:00", 7),
                    (7, "06:00", "13:00", "7,5"),
                ],
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    assert r.json()["errores"] == []
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == [
        ("2026-09-04", "06:00:00", "10:00:00", None),
        ("2026-09-04", "14:00:00", "18:00:00", None),
        ("2026-09-06", "06:00:00", "13:00:00", 420),
        ("2026-09-07", "06:00:00", "13:00:00", 450),
    ]


def test_dia_en_blanco_hereda_el_de_la_fila_anterior(client, escenario):
    """Al agregar una fila para el segundo turno del día es natural no repetir
    el número; la fila sin día es la continuación de la anterior."""
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [(4, "06:00", "10:00", None), (None, "14:00", "18:00", None)],
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == [
        ("2026-09-04", "06:00:00", "10:00:00", None),
        ("2026-09-04", "14:00:00", "18:00:00", None),
    ]


def test_reemplaza_solo_a_los_empleados_del_archivo(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana, luis = escenario["empleados"]
    for empleado in (ana, luis):
        r = client.post(
            "/turnos",
            json={
                "empleado_id": empleado["id"],
                "fecha": "2026-09-01",
                "hora_inicio": "06:00",
                "hora_fin": "18:00",
            },
        )
        assert r.status_code == 201, r.text

    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(2, "18:00", "06:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text

    # Ana queda como dice el archivo (su turno del día 1 desapareció)…
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == [
        ("2026-09-02", "18:00:00", "06:00:00", None)
    ]
    # …y Luis, que no estaba en el archivo, intacto.
    assert turnos_de(client, periodo["id"], unidad["id"], luis["id"]) == [
        ("2026-09-01", "06:00:00", "18:00:00", None)
    ]


def test_hoja_vacia_borra_los_turnos_de_ese_empleado(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    client.post(
        "/turnos",
        json={
            "empleado_id": ana["id"],
            "fecha": "2026-09-01",
            "hora_inicio": "06:00",
            "hora_fin": "18:00",
        },
    )
    contenido = construir_libro([{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": []}])
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    assert r.json()["hojas"][0]["turnos_borrados"] == 1
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == []


def test_empleado_desconocido_rechaza_todo_el_archivo(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [
            {"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]},
            {"nombre": "PEDRO NUEVO", "documento": "777", "filas": [(1, "06:00", "18:00", None)]},
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    reporte = r.json()
    assert reporte["aplicado"] is False
    assert any("777" in e["mensaje"] for e in reporte["errores"])
    # Ni siquiera los turnos de la hoja buena se guardaron.
    assert turnos_de(client, periodo["id"], unidad["id"]) == []


def test_empleado_inactivo_se_rechaza(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    assert client.patch(f"/empleados/{ana['id']}", json={"activo": False}).status_code == 200
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is False
    assert any("inactivo" in e["mensaje"] for e in r.json()["errores"])


def test_empleado_de_otra_unidad_no_sirve(client, escenario):
    periodo = escenario["periodo"]
    otra = client.post("/unidades", json={"nombre": "OTRA UNIDAD"}).json()
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], otra["id"], contenido)
    assert r.json()["aplicado"] is False
    assert any("1032456" in e["mensaje"] for e in r.json()["errores"])


def test_dia_fuera_de_la_quincena(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(20, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    reporte = r.json()
    assert reporte["aplicado"] is False
    assert reporte["errores"][0]["fila"] is not None
    assert "20" in reporte["errores"][0]["mensaje"]


def test_una_sola_hora_es_error(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", None, None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is False
    assert "salida" in r.json()["errores"][0]["mensaje"]


def test_hora_ilegible_dice_en_que_fila_esta(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "mañana", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    errores = r.json()["errores"]
    assert r.json()["aplicado"] is False
    assert errores[0]["fila"] == 8
    assert "mañana" in errores[0]["mensaje"]


def test_turnos_solapados_en_el_archivo(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [(1, "06:00", "18:00", None), (1, "14:00", "20:00", None)],
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is False
    assert "solapados" in r.json()["errores"][0]["mensaje"].lower()


def test_solapamiento_con_el_turno_nocturno_de_la_quincena_anterior(client, escenario):
    """El turno del 15 que cruza medianoche ocupa las primeras horas del 16: un
    archivo de la quincena 16–30 que empiece antes tiene que fallar."""
    unidad = escenario["unidad"]
    ana = escenario["empleados"][0]
    r = client.post(
        "/turnos",
        json={
            "empleado_id": ana["id"],
            "fecha": "2026-09-15",
            "hora_inicio": "18:00",
            "hora_fin": "06:00",
        },
    )
    assert r.status_code == 201, r.text
    segunda = client.post(
        "/periodos", json={"fecha_inicio": "2026-09-16", "fecha_fin": "2026-09-30"}
    ).json()

    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(16, "00:00", "08:00", None)]}]
    )
    r = subir(client, segunda["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is False
    assert "solapados" in r.json()["errores"][0]["mensaje"].lower()


def test_validar_solo_no_escribe_nada(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido, validar_solo=True)
    assert r.status_code == 200, r.text
    reporte = r.json()
    assert reporte["aplicado"] is False
    assert reporte["errores"] == []
    assert reporte["hojas"][0]["turnos_creados"] == 1  # lo que haría
    assert turnos_de(client, periodo["id"], unidad["id"]) == []


def test_marcas_de_quincena_se_aplican(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [(1, "06:00", "18:00", None)],
                "marcas": {"quincena_incompleta": True, "auxilio_por_dias_laborados": True},
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 200, r.text
    assert r.json()["hojas"][0]["marcas"] == {
        "quincena_incompleta": True,
        "auxilio_por_dias_laborados": True,
    }
    ajuste = client.get(
        f"/ajustes-quincena?empleado_id={ana['id']}&periodo_id={periodo['id']}"
    ).json()
    assert ajuste["quincena_incompleta"] is True
    assert ajuste["auxilio_por_dias_laborados"] is True
    assert ajuste["sin_extras"] is False


def test_sin_filas_de_marcas_no_se_pisa_lo_que_hay(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    client.put(
        f"/ajustes-quincena?empleado_id={ana['id']}&periodo_id={periodo['id']}",
        json={"quincena_incompleta": True},
    )
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    assert subir(client, periodo["id"], unidad["id"], contenido).status_code == 200
    ajuste = client.get(
        f"/ajustes-quincena?empleado_id={ana['id']}&periodo_id={periodo['id']}"
    ).json()
    assert ajuste["quincena_incompleta"] is True


def test_hoja_sin_cc_es_error_y_las_hojas_sin_tabla_se_ignoran(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [
            {
                "nombre": "ANA GOMEZ",
                "documento": "1032456",
                "filas": [(1, "06:00", "18:00", None)],
                "sin_cc": True,
            }
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    reporte = r.json()
    assert reporte["aplicado"] is False
    assert "C.C." in reporte["errores"][0]["mensaje"]
    assert reporte["ignoradas"] == ["INSTRUCCIONES"]


def test_dos_hojas_del_mismo_empleado(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [
            {"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]},
            {"nombre": "ANA (copia)", "documento": "1.032.456", "filas": [(2, "06:00", "18:00", None)]},
        ]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is False
    assert "ya venía" in r.json()["errores"][0]["mensaje"]


def test_la_cedula_se_compara_sin_puntos(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1.032.456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.json()["aplicado"] is True
    assert len(turnos_de(client, periodo["id"], unidad["id"], ana["id"])) == 1


def test_libro_sin_ninguna_hoja_de_turnos(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro([])
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 409, r.text
    assert "formato de la plantilla" in r.json()["detail"]


def test_archivo_que_no_es_excel(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    r = subir(client, periodo["id"], unidad["id"], b"esto no es un xlsx")
    assert r.status_code == 409, r.text
    assert ".xls" in r.json()["detail"]


def test_extension_equivocada_se_rechaza_antes_de_leer(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    r = client.post(
        f"/periodos/{periodo['id']}/turnos/importar?unidad_id={unidad['id']}",
        files={"archivo": ("turnos.xls", b"xxx", "application/vnd.ms-excel")},
    )
    assert r.status_code == 400, r.text
    assert ".xlsx" in r.json()["detail"]


def test_periodo_no_abierto_rechaza_la_importacion(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    # marcar el periodo como liquidado exige haber liquidado al menos una unidad
    r = client.post(f"/periodos/{periodo['id']}/liquidar", json={"unidad_id": unidad["id"]})
    assert r.status_code == 201, r.text
    assert client.post(f"/periodos/{periodo['id']}/liquidar-periodo").status_code == 200
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    r = subir(client, periodo["id"], unidad["id"], contenido)
    assert r.status_code == 409, r.text
    assert "reábralo" in r.json()["detail"]


def test_la_importacion_queda_en_la_auditoria(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    assert subir(client, periodo["id"], unidad["id"], contenido).status_code == 200
    registros = client.get("/auditoria").json()
    importaciones = [r for r in registros if r["accion"] == "importar"]
    assert len(importaciones) == 1
    assert importaciones[0]["entidad"] == "turno"


def test_el_operador_puede_importar_y_descargar_la_plantilla(client, client_operador, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    r = client_operador.get(
        f"/periodos/{periodo['id']}/turnos/plantilla?unidad_id={unidad['id']}"
    )
    assert r.status_code == 200, r.text
    contenido = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    assert subir(client_operador, periodo["id"], unidad["id"], contenido).status_code == 200


def test_plantilla_de_un_solo_empleado(client, escenario):
    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    r = client.get(
        f"/periodos/{periodo['id']}/turnos/plantilla"
        f"?unidad_id={unidad['id']}&empleado_id={ana['id']}"
    )
    assert r.status_code == 200, r.text
    assert "ANA_GOMEZ" in r.headers["content-disposition"]

    from openpyxl import load_workbook

    libro = load_workbook(BytesIO(r.content))
    assert libro.sheetnames == ["INSTRUCCIONES", "ANA GOMEZ"]


def test_archivo_demasiado_grande_se_rechaza(client, escenario):
    """El tope se aplica sin cargar el archivo entero en memoria: la ruta lee
    como máximo el límite + 1 byte."""
    from nomina.infraestructura.excel.importador import TAMANO_MAXIMO

    unidad, periodo = escenario["unidad"], escenario["periodo"]
    r = subir(client, periodo["id"], unidad["id"], b"x" * (TAMANO_MAXIMO + 10))
    assert r.status_code == 409, r.text
    assert "MB" in r.json()["detail"]


def test_la_plantilla_trae_las_horas_formuladas(client, escenario):
    """TOTAL H se calcula en Excel desde ENTRA/SALE, y el pie suma la columna:
    así la contadora cuadra las horas antes de subir el archivo.

    (Que las fórmulas den el número correcto se verifica recalculando el libro
    con LibreOffice; aquí solo se comprueba que estén y que el rango cuadre.)
    """
    from openpyxl import load_workbook

    unidad, periodo = escenario["unidad"], escenario["periodo"]
    r = client.get(f"/periodos/{periodo['id']}/turnos/plantilla?unidad_id={unidad['id']}")
    assert r.status_code == 200, r.text

    hoja = load_workbook(BytesIO(r.content))["ANA GOMEZ"]
    encabezado = next(
        f[0].row for f in hoja.iter_rows(max_col=3) if f[2].value == "ENTRA"
    )
    dias_de_la_quincena = 15  # 1 al 15 de septiembre
    primera, ultima = encabezado + 1, encabezado + dias_de_la_quincena

    for numero in range(primera, ultima + 1):
        formula = hoja.cell(row=numero, column=6).value
        assert formula.startswith("=IFERROR("), (numero, formula)
        assert f"C{numero}" in formula and f"D{numero}" in formula

    assert hoja.cell(row=ultima + 1, column=1).value == "TOTAL QUINCENA"
    assert hoja.cell(row=ultima + 1, column=6).value == f"=SUM(F{primera}:F{ultima})"


def test_celda_de_hora_con_formula_sin_calcular_se_reporta(client, escenario):
    """Una fórmula que Excel nunca calculó llega vacía al lector: sin esta
    guarda se leería como descanso y borraría los turnos del día en silencio."""
    from openpyxl import load_workbook

    unidad, periodo = escenario["unidad"], escenario["periodo"]
    ana = escenario["empleados"][0]
    r = client.post(
        "/turnos",
        json={
            "empleado_id": ana["id"],
            "fecha": "2026-09-01",
            "hora_inicio": "06:00",
            "hora_fin": "18:00",
        },
    )
    assert r.status_code == 201, r.text

    base = construir_libro(
        [{"nombre": "ANA GOMEZ", "documento": "1032456", "filas": [(1, "06:00", "18:00", None)]}]
    )
    libro = load_workbook(BytesIO(base))
    libro["ANA GOMEZ"]["C8"] = "=B8"  # fórmula sin valor cacheado
    roto = BytesIO()
    libro.save(roto)

    respuesta = subir(client, periodo["id"], unidad["id"], roto.getvalue())
    reporte = respuesta.json()
    assert reporte["aplicado"] is False
    assert reporte["errores"][0]["fila"] == 8
    assert "fórmula" in reporte["errores"][0]["mensaje"]
    # y el turno que ya estaba sigue ahí
    assert turnos_de(client, periodo["id"], unidad["id"], ana["id"]) == [
        ("2026-09-01", "06:00:00", "18:00:00", None)
    ]
