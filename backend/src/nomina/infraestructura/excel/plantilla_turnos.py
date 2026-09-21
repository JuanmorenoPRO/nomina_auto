"""Plantilla de turnos en Excel, con el formato de la tarjeta de turnos: una
hoja por empleado y una fila por día de la quincena.

Sirve para las dos direcciones: se descarga con los turnos que ya están en la
base (así es también el respaldo del cuadro de turnos) y se vuelve a subir para
cargar la quincena completa — la lee `importador.leer_plantilla_turnos`.

Nota deliberada sobre el formato de las celdas de hora: se escriben como TEXTO
(«06:00») y la columna se deja en formato General. Si se dejaran con formato de
hora, escribir `18` en Excel se guardaría como el serial 18 (18 días) y se
leería como 00:00 — un error de horas silencioso. En General, `18` queda como el
número 18 y el lector lo entiende como 18:00, `18:30` queda como hora y `6,5`
como 6.5; las tres se leen bien.

Eso es también lo que hace larga a la fórmula de TOTAL H: tiene que aceptar esos
cuatro formatos (ver `_formula_total_horas`).
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence, Set
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from nomina.dominio.entidades.empleado import Empleado
from nomina.dominio.entidades.periodo_liquidacion import PeriodoLiquidacion
from nomina.dominio.entidades.turno import TurnoRegistrado
from nomina.dominio.entidades.unidad_residencial import UnidadResidencial
from nomina.infraestructura.excel.estilos import (
    BORDE_CELDA,
    NEGRITA,
    RELLENO_ENCABEZADO,
    RELLENO_FESTIVO,
    TITULO,
    titulo_hoja,
)

DIAS_SEMANA = ["dom", "lun", "mar", "mié", "jue", "vie", "sáb"]
MARCA_FESTIVO = "✦"

COLUMNAS = ["DÍA", "N.º", "ENTRA", "SALE", "JORN. ORD. (h)", "TOTAL H"]

_ETIQUETAS_MARCAS = [
    ("quincena_incompleta", "QUINCENA INCOMPLETA:"),
    ("sin_extras", "NO CALCULAR EXTRAS:"),
    ("auxilio_por_dias_laborados", "AUXILIO PRORRATEADO:"),
    ("pagar_dia_31", "PAGAR DÍA 31:"),
]


def _dias(periodo: PeriodoLiquidacion) -> list[date]:
    dias = []
    dia = periodo.fecha_inicio
    while dia <= periodo.fecha_fin:
        dias.append(dia)
        dia += timedelta(days=1)
    return dias


def _abreviatura(dia: date) -> str:
    return DIAS_SEMANA[dia.isoweekday() % 7]


def _horas(minutos: int) -> float | int:
    horas = Decimal(minutos) / 60
    return int(horas) if horas == horas.to_integral_value() else float(round(horas, 2))


def _hhmm(hora: object) -> str:
    return f"{hora:%H:%M}"  # type: ignore[str-format]


def _como_fraccion_de_dia(ref: str) -> str:
    """Normaliza una celda de hora a fracción de día, sea lo que sea que haya
    digitado la contadora. Excel guarda cuatro cosas distintas en esa celda:

    - `06:00` escrito por la plantilla → texto      → `TIMEVALUE` → 0.25
    - `18:30` digitado                → hora real   → número < 1, ya es fracción
    - `18` digitado                   → número 18   → son horas: `/24`
    - `6,5` digitado                  → número 6.5  → son horas: `/24`

    El `&":00"` es para un `"6"` que haya quedado como texto.
    """
    return (
        f"IF(ISNUMBER({ref}),IF({ref}>=1,{ref}/24,{ref}),"
        f'TIMEVALUE(IF(ISERROR(FIND(":",{ref})),{ref}&":00",{ref})))'
    )


def _formula_total_horas(fila: int) -> str:
    """Horas del turno de esa fila, calculadas por Excel.

    `MOD(...,1)` es lo que resuelve el cruce de medianoche (18:00→06:00 son 12 h,
    no −12), y entrada igual a salida son 24 h: la misma semántica que
    `Turno.intervalo()`. Sin horario la celda queda vacía (descanso, no cero) y
    con una hora ilegible muestra «?» en vez de inventar un número.
    """
    ini = _como_fraccion_de_dia(f"C{fila}")
    fin = _como_fraccion_de_dia(f"D{fila}")
    return (
        f'=IFERROR(IF(OR(C{fila}="",D{fila}=""),"",'
        f'IF({ini}={fin},24,MOD({fin}-{ini},1)*24)),"?")'
    )


def _hoja_empleado(
    hoja: Worksheet,
    unidad: UnidadResidencial,
    periodo: PeriodoLiquidacion,
    empleado: Empleado,
    turnos: Sequence[TurnoRegistrado],
    festivos: Set[date],
    marcas: Mapping[str, bool],
) -> None:
    hoja["A1"] = "EDIFICIO:"
    hoja["B1"] = unidad.nombre
    hoja["A2"] = "NOMBRE Y APELLIDO:"
    hoja["B2"] = empleado.nombre
    hoja["A3"] = "C.C.:"
    celda_documento = hoja["B3"]
    celda_documento.value = empleado.documento
    celda_documento.number_format = "@"  # texto: no perder ceros a la izquierda
    hoja["A4"] = "MES:"
    hoja["B4"] = (
        f"{periodo.fecha_inicio:%d/%m/%Y} al {periodo.fecha_fin:%d/%m/%Y}"
    )
    for fila in range(1, 5):
        hoja.cell(row=fila, column=1).font = NEGRITA
    hoja["B2"].font = NEGRITA

    fila = 6
    dias = _dias(periodo)
    hay_dia_31 = any(d.day >= 31 for d in dias)
    for clave, etiqueta in _ETIQUETAS_MARCAS:
        if clave == "pagar_dia_31" and not hay_dia_31:
            continue
        hoja.cell(row=fila, column=1, value=etiqueta)
        hoja.cell(row=fila, column=2, value="SÍ" if marcas.get(clave) else "NO")
        fila += 1

    fila += 1
    fila_encabezado = fila
    for columna, texto in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=fila_encabezado, column=columna, value=texto)
        celda.font = NEGRITA
        celda.fill = RELLENO_ENCABEZADO
        celda.border = BORDE_CELDA
        celda.alignment = Alignment(horizontal="center")

    por_dia: dict[date, list[TurnoRegistrado]] = {}
    for turno in turnos:
        por_dia.setdefault(turno.turno.fecha, []).append(turno)

    fila = fila_encabezado
    for dia in dias:
        es_festivo = dia in festivos
        # Rosado para domingos y festivos; el ✦ solo para los festivos, igual
        # que en el cuadro de turnos de la UI.
        especial = dia.isoweekday() == 7 or es_festivo
        del_dia = sorted(por_dia.get(dia, []), key=lambda t: t.turno.hora_inicio)
        # Un día sin turnos también lleva su fila: es el descanso, y así la
        # contadora escribe el horario sobre la fila que ya existe.
        for indice in range(max(1, len(del_dia))):
            fila += 1
            turno = del_dia[indice].turno if indice < len(del_dia) else None
            etiqueta_dia = _abreviatura(dia) + (f" {MARCA_FESTIVO}" if es_festivo else "")
            hoja.cell(row=fila, column=1, value=etiqueta_dia)
            hoja.cell(row=fila, column=2, value=dia.day)
            # La fórmula va en TODAS las filas, también en las de descanso: así
            # el total aparece solo en cuanto se le escriba un horario.
            hoja.cell(row=fila, column=6, value=_formula_total_horas(fila))
            if turno is not None:
                hoja.cell(row=fila, column=3, value=_hhmm(turno.hora_inicio))
                hoja.cell(row=fila, column=4, value=_hhmm(turno.hora_fin))
                if turno.minutos_jornada_ordinaria is not None:
                    hoja.cell(
                        row=fila, column=5, value=_horas(turno.minutos_jornada_ordinaria)
                    )
            for columna in range(1, len(COLUMNAS) + 1):
                celda = hoja.cell(row=fila, column=columna)
                celda.border = BORDE_CELDA
                celda.alignment = Alignment(horizontal="center")
                if especial:
                    celda.fill = RELLENO_FESTIVO

    ultima_fila_datos = fila
    fila += 1
    hoja.cell(row=fila, column=1, value="TOTAL QUINCENA").font = NEGRITA
    celda_total = hoja.cell(
        row=fila,
        column=6,
        value=f"=SUM(F{fila_encabezado + 1}:F{ultima_fila_datos})",
    )
    celda_total.font = NEGRITA
    celda_total.border = BORDE_CELDA

    hoja.column_dimensions["A"].width = 20
    hoja.column_dimensions["B"].width = 10
    for letra in ("C", "D"):
        hoja.column_dimensions[letra].width = 12
    hoja.column_dimensions["E"].width = 16
    hoja.column_dimensions["F"].width = 12
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)


_INSTRUCCIONES = [
    ("Cómo llenar esta plantilla", TITULO),
    ("", None),
    ("Hay una hoja por empleado, con una fila por día de la quincena.", None),
    ("", None),
    ("• ENTRA y SALE: el horario del día. Se acepta 06:00, 6, 18:30 o 6,5.", None),
    ("  Si la salida es menor o igual que la entrada, el turno cruza la medianoche", None),
    ("  (18:00 a 06:00 son 12 horas del día siguiente).", None),
    ("• Día sin horario = descanso. No hay que escribir nada.", None),
    ("• Turno partido: agregue otra fila con el mismo número de día.", None),
    ("• JORN. ORD. (h): marca el turno como «jornada ordinaria» y dice cuántas", None),
    ("  horas cubre el salario. Solo el excedente se paga, como hora extra.", None),
    ("  Déjela vacía en los turnos normales.", None),
    ("• Las filas SÍ/NO del encabezado son las marcas de la quincena. Si borra la", None),
    ("  fila completa, la marca que haya en el sistema no se modifica.", None),
    ("", None),
    ("Al importar:", NEGRITA),
    ("• El empleado se identifica por la C.C. del encabezado, NO por el nombre de", None),
    ("  la hoja. No borre esa fila ni cambie la cédula.", None),
    ("• Los turnos de la quincena de cada empleado que aparezca en el archivo se", None),
    ("  REEMPLAZAN por lo que diga su hoja. Los empleados que no estén en el", None),
    ("  archivo no se tocan.", None),
    ("• La importación no crea empleados: si una C.C. no existe en la unidad, el", None),
    ("  archivo se rechaza completo y no se guarda nada.", None),
    ("• TOTAL H y TOTAL QUINCENA se calculan solos: no escriba encima de esas", None),
    ("  celdas. Si agrega una fila para un turno partido, copie la fórmula de", None),
    ("  TOTAL H de la fila de arriba. Al importar se recalculan de todos modos.", None),
    ("• Puede agregar o quitar filas y columnas; el encabezado de la tabla se", None),
    ("  reconoce por los títulos ENTRA y SALE.", None),
    ("• Guarde el archivo como .xlsx (no .xls).", None),
]


def _hoja_instrucciones(hoja: Worksheet, unidad: UnidadResidencial, periodo: PeriodoLiquidacion) -> None:
    hoja["A1"] = f"{unidad.nombre} — cuadro de turnos"
    hoja["A1"].font = TITULO
    hoja["A2"] = f"Quincena del {periodo.fecha_inicio:%d/%m/%Y} al {periodo.fecha_fin:%d/%m/%Y}"
    hoja["A2"].font = NEGRITA
    for desplazamiento, (texto, fuente) in enumerate(_INSTRUCCIONES, start=4):
        celda = hoja.cell(row=desplazamiento, column=1, value=texto)
        if fuente is not None:
            celda.font = fuente
    hoja.column_dimensions["A"].width = 84


def generar_plantilla_turnos(
    unidad: UnidadResidencial,
    periodo: PeriodoLiquidacion,
    empleados: Sequence[Empleado],
    turnos: Mapping[UUID, Sequence[TurnoRegistrado]],
    festivos: Set[date],
    marcas: Mapping[UUID, Mapping[str, bool]] | None = None,
) -> bytes:
    libro = Workbook()
    instrucciones = libro.active
    assert instrucciones is not None
    instrucciones.title = "INSTRUCCIONES"
    _hoja_instrucciones(instrucciones, unidad, periodo)

    usados = {"INSTRUCCIONES"}
    for empleado in sorted(empleados, key=lambda e: e.nombre):
        hoja = libro.create_sheet(titulo_hoja(empleado.nombre, usados))
        _hoja_empleado(
            hoja,
            unidad,
            periodo,
            empleado,
            turnos.get(empleado.id, ()),
            festivos,
            (marcas or {}).get(empleado.id, {}),
        )

    contenido = BytesIO()
    libro.save(contenido)
    return contenido.getvalue()


def nombre_archivo_plantilla(
    unidad: UnidadResidencial, periodo: PeriodoLiquidacion, empleado: Empleado | None = None
) -> str:
    def limpiar(texto: str) -> str:
        return "".join(c if c.isalnum() else "_" for c in texto).strip("_")

    quien = f"_{limpiar(empleado.nombre)}" if empleado is not None else ""
    return (
        f"turnos_{limpiar(unidad.nombre)}{quien}"
        f"_{periodo.fecha_inicio:%Y-%m-%d}_{periodo.fecha_fin:%Y-%m-%d}.xlsx"
    )
