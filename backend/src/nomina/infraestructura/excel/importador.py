"""Lector de la plantilla de turnos en Excel (formato tarjeta: una hoja por
empleado, una fila por día).

Es deliberadamente tolerante, porque la contadora va a mover filas, renombrar
hojas y escribir las horas como se le dé la gana:

- la fila de encabezado se busca (la primera con ENTRA y SALE) en vez de
  asumirse en una posición fija, y las columnas se resuelven por nombre;
- las etiquetas del encabezado se comparan sin acentos, puntos ni espacios,
  así que «C.C.», «CC:» y «c c» son lo mismo;
- las horas se aceptan como las devuelve Excel (`time`, `datetime`,
  `timedelta`) y como texto o número (`"18:30"`, `18`, `18.5`, `"6,5"`).

Solo traduce el archivo a datos puros; las reglas de negocio están en
`aplicacion/casos_uso/importar_turnos.py`.
"""

from __future__ import annotations

import unicodedata
from datetime import datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from nomina.aplicacion.casos_uso.importar_turnos import (
    FilaTurno,
    HojaTurnos,
    LibroTurnos,
    MarcasHoja,
)

TAMANO_MAXIMO = 10 * 1024 * 1024  # 10 MB: una plantilla real pesa unos pocos KB
MAXIMO_HOJAS = 200
MAXIMO_FILAS = 400  # una quincena son ~16 filas; el resto es margen para notas

_COLUMNAS_DIA = {"NO", "N", "NRO", "NUM", "NUMERO", "DIADELMES"}
_COLUMNAS_ENTRA = {"ENTRA", "ENTRADA", "INICIO", "HORAENTRADA", "HORAINICIO"}
_COLUMNAS_SALE = {"SALE", "SALIDA", "FIN", "HORASALIDA", "HORAFIN"}
_COLUMNAS_JORNADA = {
    "JORNORD",
    "JORNORDH",
    "JORNADAORD",
    "JORNADAORDH",
    "JORNADAORDINARIA",
    "JORNADAORDINARIAH",
    "ORD",
}

_ETIQUETA_DOCUMENTO = {"CC", "CEDULA", "DOCUMENTO", "NODOCUMENTO", "IDENTIFICACION"}
_ETIQUETA_NOMBRE = {"NOMBREYAPELLIDO", "NOMBRE", "NOMBREYAPELLIDOS", "EMPLEADO"}
_ETIQUETA_UNIDAD = {"EDIFICIO", "UNIDAD", "UNIDADRESIDENCIAL"}

_ETIQUETAS_MARCAS = {
    "quincena_incompleta": {"QUINCENAINCOMPLETA", "NOLABOROTODALAQUINCENA"},
    "sin_extras": {"NOCALCULAREXTRAS", "NOCALCULARHORASEXTRA", "SINEXTRAS"},
    "auxilio_por_dias_laborados": {"AUXILIOPRORRATEADO", "AUXILIOPORDIASLABORADOS"},
    "pagar_dia_31": {"PAGARDIA31", "DIA31"},
}

_VERDADEROS = {"SI", "S", "X", "TRUE", "VERDADERO", "1", "SÍ"}
_FALSOS = {"NO", "N", "FALSE", "FALSO", "0"}


def normalizar(texto: Any) -> str:
    """Sin acentos, sin nada que no sea alfanumérico, en mayúsculas:
    «Jorn. ord. (h)» → «JORNORDH»."""
    if texto is None:
        return ""
    plano = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in plano if c.isalnum() and not unicodedata.combining(c)).upper()


def _leer_hora(valor: Any) -> time | None:
    """`None` si la celda está vacía; `ValueError` si no es una hora."""
    if valor is None:
        return None
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, timedelta):
        minutos = int(valor.total_seconds() // 60) % (24 * 60)
        return time(minutos // 60, minutos % 60)
    if isinstance(valor, bool):
        raise ValueError(f"«{valor}» no es una hora")
    if isinstance(valor, (int, float, Decimal)):
        return _hora_de_numero(Decimal(str(valor)), valor)
    texto = str(valor).strip()
    if not texto:
        return None
    partes = texto.split(":")
    if len(partes) > 1:
        try:
            horas = int(partes[0])
            minutos = int(partes[1])
        except ValueError:
            raise ValueError(f"«{texto}» no es una hora (se espera algo como 18:30)") from None
        if not (0 <= horas <= 23 and 0 <= minutos <= 59):
            raise ValueError(f"«{texto}» no es una hora válida")
        return time(horas, minutos)
    try:
        numero = Decimal(texto.replace(",", "."))
    except InvalidOperation:
        raise ValueError(f"«{texto}» no es una hora (se espera algo como 18:30 o 18)") from None
    return _hora_de_numero(numero, texto)


def _hora_de_numero(numero: Decimal, original: Any) -> time:
    """Un número es la hora del día: 6 → 06:00, 18.5 → 18:30."""
    if not (0 <= numero < 24):
        raise ValueError(f"«{original}» no es una hora del día (0 a 23:59)")
    minutos_exactos = (numero * 60).quantize(Decimal("0.0001"))
    if minutos_exactos != minutos_exactos.to_integral_value():
        raise ValueError(f"«{original}» no cae en un minuto exacto")
    minutos = int(minutos_exactos)
    return time(minutos // 60, minutos % 60)


def _leer_dia(valor: Any) -> int | None:
    """El número del día; tolera el «6✦» con el que se marcan los festivos."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, datetime):
        return valor.day
    if isinstance(valor, (int, float, Decimal)):
        return int(valor)
    digitos = ""
    for caracter in str(valor).strip():
        if caracter.isdigit():
            digitos += caracter
        elif digitos:
            break
    return int(digitos) if digitos else None


def _leer_horas_jornada(valor: Any) -> Decimal | None:
    """Horas de jornada ordinaria: acepta 7, «7,5» y lo que Excel devuelva."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        raise ValueError("La jornada ordinaria se escribe en horas, no como sí/no")
    if isinstance(valor, time):
        return Decimal(valor.hour) + Decimal(valor.minute) / 60
    if isinstance(valor, timedelta):
        return Decimal(int(valor.total_seconds() // 60)) / 60
    if isinstance(valor, (int, float, Decimal)):
        horas = Decimal(str(valor))
    else:
        texto = str(valor).strip()
        if not texto:
            return None
        if ":" in texto:
            hora = _leer_hora(texto)
            assert hora is not None
            return Decimal(hora.hour) + Decimal(hora.minute) / 60
        try:
            horas = Decimal(texto.replace(",", "."))
        except InvalidOperation:
            raise ValueError(f"«{texto}» no es un número de horas") from None
    if horas <= 0 or horas > 24:
        raise ValueError(f"La jornada ordinaria debe estar entre 0 y 24 horas, no {horas}")
    return horas


def _leer_marca(valor: Any) -> bool | None:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return valor
    texto = str(valor).strip()
    if not texto:
        return None
    normalizado = normalizar(texto)
    if normalizado in {normalizar(v) for v in _VERDADEROS}:
        return True
    if normalizado in {normalizar(v) for v in _FALSOS}:
        return False
    return None


def _fila_encabezado(filas: list[tuple[Any, ...]]) -> int | None:
    for indice, fila in enumerate(filas):
        normalizadas = {normalizar(c) for c in fila}
        if normalizadas & _COLUMNAS_ENTRA and normalizadas & _COLUMNAS_SALE:
            return indice
    return None


def _mapa_columnas(fila: tuple[Any, ...]) -> dict[str, int]:
    mapa: dict[str, int] = {}
    for columna, celda in enumerate(fila):
        nombre = normalizar(celda)
        if not nombre:
            continue
        if nombre in _COLUMNAS_ENTRA and "entra" not in mapa:
            mapa["entra"] = columna
        elif nombre in _COLUMNAS_SALE and "sale" not in mapa:
            mapa["sale"] = columna
        elif nombre in _COLUMNAS_DIA and "dia" not in mapa:
            mapa["dia"] = columna
        elif nombre in _COLUMNAS_JORNADA and "jornada" not in mapa:
            mapa["jornada"] = columna
    return mapa


def _celda(fila: tuple[Any, ...], columna: int | None) -> Any:
    if columna is None or columna >= len(fila):
        return None
    valor = fila[columna]
    if isinstance(valor, str) and not valor.strip():
        return None
    return valor


def _leer_encabezado(filas: list[tuple[Any, ...]]) -> tuple[dict[str, str], MarcasHoja]:
    """Etiqueta → valor, buscando en cada fila la primera celda con texto (la
    etiqueta) y la siguiente con contenido (el valor)."""
    datos: dict[str, str] = {}
    marcas: dict[str, bool | None] = {}
    for fila in filas:
        etiqueta = ""
        valor: Any = None
        for celda in fila:
            if celda is None or (isinstance(celda, str) and not celda.strip()):
                continue
            if not etiqueta:
                etiqueta = normalizar(celda)
            else:
                valor = celda
                break
        if not etiqueta:
            continue
        if etiqueta in _ETIQUETA_DOCUMENTO:
            datos.setdefault("documento", _texto_documento(valor))
        elif etiqueta in _ETIQUETA_NOMBRE:
            datos.setdefault("nombre_empleado", str(valor).strip() if valor else "")
        elif etiqueta in _ETIQUETA_UNIDAD:
            datos.setdefault("unidad", str(valor).strip() if valor else "")
        else:
            for marca, alias in _ETIQUETAS_MARCAS.items():
                if etiqueta in alias and marca not in marcas:
                    marcas[marca] = _leer_marca(valor)
    return datos, MarcasHoja(**marcas)


def _texto_documento(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _leer_hoja(nombre: str, filas: list[tuple[Any, ...]]) -> HojaTurnos | None:
    indice = _fila_encabezado(filas)
    if indice is None:
        return None  # no es una hoja de turnos (INSTRUCCIONES, RESUMEN, notas…)

    columnas = _mapa_columnas(filas[indice])
    datos, marcas = _leer_encabezado(filas[:indice])
    errores: list[str] = []
    documento = datos.get("documento", "")
    if not documento:
        errores.append(
            "La hoja no dice la C.C. del empleado: agregue la fila «C.C.:» en el "
            "encabezado (es lo único con lo que se identifica al empleado)."
        )

    filas_turno: list[FilaTurno] = []
    ultimo_dia: int | None = None
    for desplazamiento, fila in enumerate(filas[indice + 1 :], start=indice + 2):
        primera = normalizar(_celda(fila, 0))
        if primera.startswith("TOTAL"):
            break  # pie de tabla («Total quincena»)

        dia = _leer_dia(_celda(fila, columnas.get("dia")))
        if dia is not None:
            ultimo_dia = dia

        bruto_entra = _celda(fila, columnas.get("entra"))
        bruto_sale = _celda(fila, columnas.get("sale"))
        bruto_jornada = _celda(fila, columnas.get("jornada"))
        if bruto_entra is None and bruto_sale is None:
            continue  # descanso

        error = ""
        entra = sale = None
        horas_jornada = None
        try:
            entra = _leer_hora(bruto_entra)
            sale = _leer_hora(bruto_sale)
            horas_jornada = _leer_horas_jornada(bruto_jornada)
        except ValueError as e:
            error = str(e)

        filas_turno.append(
            FilaTurno(
                fila=desplazamiento,
                dia_del_mes=dia if dia is not None else ultimo_dia,
                hora_inicio=entra,
                hora_fin=sale,
                horas_jornada_ordinaria=horas_jornada,
                error=error,
            )
        )

    return HojaTurnos(
        nombre_hoja=nombre,
        documento=documento,
        nombre_empleado=datos.get("nombre_empleado", ""),
        unidad=datos.get("unidad", ""),
        filas=tuple(filas_turno),
        marcas=marcas,
        errores=tuple(errores),
    )


def leer_plantilla_turnos(contenido: bytes) -> LibroTurnos:
    """Traduce el .xlsx a `LibroTurnos`. `ValueError` si no es un Excel legible."""
    if not contenido:
        raise ValueError("El archivo está vacío")
    if len(contenido) > TAMANO_MAXIMO:
        raise ValueError(
            f"El archivo pesa más de {TAMANO_MAXIMO // (1024 * 1024)} MB; "
            "la plantilla de turnos pesa unos pocos KB"
        )

    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # openpyxl lanza de todo: BadZipFile, KeyError, InvalidFile…
        raise ValueError(
            "No se pudo leer el archivo como Excel (.xlsx). Si es un .xls antiguo, "
            "ábralo en Excel y guárdelo como .xlsx."
        ) from e

    try:
        if len(libro.sheetnames) > MAXIMO_HOJAS:
            raise ValueError(f"El archivo tiene más de {MAXIMO_HOJAS} hojas")

        hojas: list[HojaTurnos] = []
        ignoradas: list[str] = []
        for nombre in libro.sheetnames:
            pagina = libro[nombre]
            filas = [
                fila
                for _, fila in zip(
                    range(MAXIMO_FILAS), pagina.iter_rows(values_only=True), strict=False
                )
            ]
            hoja = _leer_hoja(nombre, filas)
            if hoja is None:
                ignoradas.append(nombre)
            else:
                hojas.append(hoja)
        return LibroTurnos(hojas=tuple(hojas), ignoradas=tuple(ignoradas))
    finally:
        libro.close()
