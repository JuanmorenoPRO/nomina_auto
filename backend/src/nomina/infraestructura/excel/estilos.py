"""Estilos y helpers compartidos por los libros de Excel que genera la app
(la liquidación en `exportador.py` y la plantilla de turnos en
`plantilla_turnos.py`), para que ambos se vean como la planilla de la contadora."""

from __future__ import annotations

import re

from openpyxl.styles import Border, Font, PatternFill, Side

NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=13)
PESOS = "#,##0"
BORDE_FINO = Border(bottom=Side(style="thin"))

# Domingos y festivos van rosados, igual que en el cuadro de turnos de la UI.
RELLENO_FESTIVO = PatternFill("solid", fgColor="FCE4EC")
RELLENO_ENCABEZADO = PatternFill("solid", fgColor="EEEEEE")

BORDE_CELDA = Border(
    left=Side(style="thin", color="BBBBBB"),
    right=Side(style="thin", color="BBBBBB"),
    top=Side(style="thin", color="BBBBBB"),
    bottom=Side(style="thin", color="BBBBBB"),
)


def titulo_hoja(nombre: str, usados: set[str] | None = None) -> str:
    """Excel limita títulos a 31 caracteres y prohíbe []:*?/\\

    Con `usados` se desambigua: dos empleados cuyo nombre se trunca igual no
    pueden compartir el título de la hoja.
    """
    limpio = re.sub(r"[\[\]:*?/\\]", "", nombre).strip()[:31] or "EMPLEADO"
    if usados is None:
        return limpio
    candidato = limpio
    sufijo = 2
    while candidato in usados:
        marca = f" ({sufijo})"
        candidato = limpio[: 31 - len(marca)] + marca
        sufijo += 1
    usados.add(candidato)
    return candidato
